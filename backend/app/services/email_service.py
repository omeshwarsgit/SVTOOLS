import os
import io
import smtplib
from datetime import datetime, date
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.application import MIMEApplication
from pathlib import Path
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.models.schemas import MISDashboardData, MISTabItem


def generate_styled_excel_report(mis_data: MISDashboardData) -> bytes:
    """
    Generate an executive multi-tab Excel workbook (.xlsx) containing:
    1. Executive Summary (KPI metrics, discrepancy category breakdown, channel distribution)
    2. Cancellation Pending
    3. Confirmed SU / Cancelled PMS
    4. Missing in PMS
    5. Missing in SU
    6. PMS Special Status (Tentative / No-show)
    7. Base vs Query Mismatch
    8. Query not in Base
    9. All SU Bookings
    """
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling constants
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    alert_fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
    card_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    # -------------------------------------------------------------
    # TAB 1: EXECUTIVE SUMMARY
    # -------------------------------------------------------------
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_sum.merge_cells("A1:G1")
    t_cell = ws_sum["A1"]
    t_cell.value = "STAYVISTA OPERATIONAL RECONCILIATION & MIS EXECUTIVE SUMMARY"
    t_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    t_cell.fill = header_fill
    t_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 35

    # Run Metadata
    ws_sum["A3"] = "Generated At:"
    ws_sum["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_sum["A4"] = "Batch ID:"
    ws_sum["B4"] = mis_data.batch_id
    ws_sum["A3"].font = Font(bold=True)
    ws_sum["A4"].font = Font(bold=True)

    # Primary Metrics Table
    ws_sum["A6"] = "Primary Reconciliation Metrics"
    ws_sum["A6"].font = Font(bold=True, size=12, color="1E3A8A")
    
    primary_headers = ["Metric", "Count", "Description"]
    for col_idx, h in enumerate(primary_headers, 1):
        c = ws_sum.cell(row=7, column=col_idx, value=h)
        c.font = header_font
        c.fill = accent_fill
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    p_metrics = [
        ("Total SU Bookings Ingested", mis_data.primary_metrics.total_bookings, "Total records received in Source of Truth (SU)"),
        ("Matched Bookings", mis_data.primary_metrics.matched, "Successfully cross-referenced and confirmed in PMS"),
        ("Mismatched Bookings", mis_data.primary_metrics.mismatched, "Status discrepancies (e.g. SU Confirmed vs PMS Cancelled)"),
        ("Missing in PMS", mis_data.primary_metrics.missing, "Valid SU bookings absent in PMS property dump"),
        ("Cancellation Pending (CRITICAL)", mis_data.primary_metrics.cancellation_pending, "In-transit cancellations requiring immediate OTA action"),
    ]

    for row_idx, (label, val, desc) in enumerate(p_metrics, 8):
        c1 = ws_sum.cell(row=row_idx, column=1, value=label)
        c2 = ws_sum.cell(row=row_idx, column=2, value=val)
        c3 = ws_sum.cell(row=row_idx, column=3, value=desc)
        for c in (c1, c2, c3):
            c.border = thin_border
        c2.alignment = Alignment(horizontal="right")
        c2.font = Font(bold=True)
        if "CRITICAL" in label:
            c1.font = Font(bold=True, color="DC2626")
            c2.font = Font(bold=True, color="DC2626")

    # Discrepancy Breakdown by Tab
    start_tab_row = 15
    ws_sum.cell(row=start_tab_row, column=1, value="Discrepancy Category Breakdown").font = Font(bold=True, size=12, color="1E3A8A")
    tab_headers = ["Category / Tab", "Actionable Records", "Operational Target"]
    for col_idx, h in enumerate(tab_headers, 1):
        c = ws_sum.cell(row=start_tab_row + 1, column=col_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")
        c.border = thin_border

    tab_descriptions = {
        "cancellation_pending": ("Cancellation Pending", "Cancel in PMS / Notify OTA Partner"),
        "confirmed_su_cancelled_pms": ("Confirmed in SU / Cancelled in PMS", "Verify guest status & channel commission"),
        "missing_in_pms": ("Missing in PMS", "Inject booking into PMS or sync OTA extranet"),
        "missing_in_su": ("Missing in SU (PMS-only)", "Audit manual PMS entries or missing SU feed"),
        "pms_special_status": ("PMS Special Status (Tentative/No-show)", "Follow up with Relationship Manager"),
        "base_vs_query_mismatch": ("Base vs Query Dump Mismatch", "Sync PMS database replicas"),
        "query_not_in_base": ("Query records missing in Base", "Resolve PMS audit query discrepancies"),
        "all_su_bookings": ("All Ingested SU Bookings", "Master record archive"),
    }

    cur_row = start_tab_row + 2
    for tab_key, (tab_title, op_target) in tab_descriptions.items():
        count = mis_data.tab_counts.get(tab_key, len(mis_data.tabs_data.get(tab_key, [])))
        c1 = ws_sum.cell(row=cur_row, column=1, value=tab_title)
        c2 = ws_sum.cell(row=cur_row, column=2, value=count)
        c3 = ws_sum.cell(row=cur_row, column=3, value=op_target)
        for c in (c1, c2, c3):
            c.border = thin_border
        c2.alignment = Alignment(horizontal="right")
        c2.font = Font(bold=True)
        cur_row += 1

    # Secondary PMS Audit Metrics
    cur_row += 1
    ws_sum.cell(row=cur_row, column=1, value="Secondary PMS Audit Statistics").font = Font(bold=True, size=12, color="1E3A8A")
    cur_row += 1
    sec_metrics = [
        ("PMS Base Dump Rows", mis_data.secondary_metrics.pms_base_rows),
        ("PMS Query Dump Rows", mis_data.secondary_metrics.pms_query_rows),
        ("PMS Tentative / No-show Rows", mis_data.secondary_metrics.pms_tentative_noshow),
        ("PMS Not Found in SU", mis_data.secondary_metrics.pms_not_found_in_su),
        ("Base vs Query Mismatch", mis_data.secondary_metrics.base_vs_query_mismatch),
        ("Query Missing in Base", mis_data.secondary_metrics.query_missing_in_base),
    ]
    for m_label, m_val in sec_metrics:
        c1 = ws_sum.cell(row=cur_row, column=1, value=m_label)
        c2 = ws_sum.cell(row=cur_row, column=2, value=m_val)
        c1.border = thin_border
        c2.border = thin_border
        c2.alignment = Alignment(horizontal="right")
        cur_row += 1

    # Auto-adjust summary column widths
    ws_sum.column_dimensions["A"].width = 38
    ws_sum.column_dimensions["B"].width = 22
    ws_sum.column_dimensions["C"].width = 50

    # -------------------------------------------------------------
    # TAB 2 - 9: OPERATIONAL DISCREPANCY TABS
    # -------------------------------------------------------------
    tab_configs = [
        ("cancellation_pending", "Cancellation Pending", alert_fill),
        ("confirmed_su_cancelled_pms", "SU Confirmed - PMS Cancelled", header_fill),
        ("missing_in_pms", "Missing in PMS", header_fill),
        ("missing_in_su", "Missing in SU", header_fill),
        ("pms_special_status", "PMS Special Status", header_fill),
        ("base_vs_query_mismatch", "Base vs Query Mismatch", header_fill),
        ("query_not_in_base", "Query not in Base", header_fill),
        ("all_su_bookings", "All SU Bookings", accent_fill),
    ]

    detail_headers = [
        "Reservation ID",
        "Vendor Booking ID",
        "Guest Name",
        "Channel",
        "SU Status",
        "PMS Status",
        "Property ID",
        "Property Name",
        "Assigned RM",
        "Target Portal Link",
        "Check-In",
        "Check-Out",
        "Match Type",
        "Notes"
    ]

    for tab_key, tab_title, fill_style in tab_configs:
        items: List[MISTabItem] = mis_data.tabs_data.get(tab_key, [])
        # Excel sheet name limit is 31 chars
        clean_sheet_title = tab_title[:31].replace("/", "-")
        ws = wb.create_sheet(title=clean_sheet_title)
        ws.views.sheetView[0].showGridLines = True

        # Header Row
        for col_idx, h in enumerate(detail_headers, 1):
            c = ws.cell(row=1, column=col_idx, value=h)
            c.font = header_font
            c.fill = fill_style
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = thin_border
        ws.row_dimensions[1].height = 28

        # Data Rows
        for r_idx, it in enumerate(items, 2):
            row_vals = [
                it.reservation_id,
                it.vendor_booking_id or "",
                it.guest_name or "",
                it.channel,
                it.su_status or "",
                it.pms_status or "",
                it.property_id or "",
                it.property_name or "",
                it.assigned_representative or "Unassigned",
                it.target_portal_url or "",
                it.check_in or "",
                it.check_out or "",
                it.match_type or "",
                it.notes or ""
            ]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r_idx, column=col_idx, value=val)
                cell.border = thin_border
                if col_idx in (1, 2, 7, 11, 12):
                    cell.alignment = Alignment(horizontal="center")
                elif col_idx in (4, 5, 6):
                    cell.alignment = Alignment(horizontal="center")
                else:
                    cell.alignment = Alignment(horizontal="left")

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[:100]:  # check first 100 rows for speed
                v = str(cell.value or "")
                if len(v) > max_len:
                    max_len = len(v)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_email_html_body(mis_data: MISDashboardData) -> str:
    """
    Build a modern, executive HTML email body with responsive layout,
    KPI badges, category breakdown table, and high-priority action items.
    """
    today_str = date.today().strftime("%d %B %Y")
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    p = mis_data.primary_metrics
    pending_count = p.cancellation_pending
    pending_items = mis_data.tabs_data.get("cancellation_pending", [])[:10]  # top 10

    # Build action items rows
    action_rows_html = ""
    for idx, it in enumerate(pending_items, 1):
        url_link = (
            f'<a href="{it.target_portal_url}" target="_blank" style="color: #2563EB; text-decoration: underline;">Portal Link</a>'
            if it.target_portal_url and it.target_portal_url != "not available"
            else '<span style="color: #94A3B8;">N/A</span>'
        )
        action_rows_html += f"""
        <tr style="border-bottom: 1px solid #E2E8F0; font-size: 13px;">
            <td style="padding: 10px; font-weight: 600; color: #1E293B;">{it.reservation_id}</td>
            <td style="padding: 10px; color: #475569;">{it.vendor_booking_id or '-'}</td>
            <td style="padding: 10px; color: #1E293B;">{it.guest_name or 'Unknown'}</td>
            <td style="padding: 10px;"><span style="background-color: #EFF6FF; color: #1D4ED8; padding: 3px 8px; border-radius: 4px; font-weight: 500; font-size: 12px;">{it.channel}</span></td>
            <td style="padding: 10px; color: #475569;">{it.property_name or f'ID #{it.property_id}'}</td>
            <td style="padding: 10px; color: #0F766E; font-weight: 500;">{it.assigned_representative or 'Unassigned'}</td>
            <td style="padding: 10px; color: #64748B;">{it.check_in or '-'}</td>
            <td style="padding: 10px;">{url_link}</td>
        </tr>
        """

    if not action_rows_html:
        action_rows_html = """
        <tr>
            <td colspan="8" style="padding: 16px; text-align: center; color: #10B981; font-weight: 600;">
                No pending cancellations found. All records synchronized!
            </td>
        </tr>
        """

    html = f"""
    <!DOCTYPE html>
    <html>
    <head>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <title>StayVista MIS Reconciliation Report</title>
    </head>
    <body style="font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif; background-color: #F8FAFC; margin: 0; padding: 24px; color: #1E293B;">
        <table align="center" border="0" cellpadding="0" cellspacing="0" width="100%" style="max-width: 820px; background-color: #FFFFFF; border-radius: 8px; overflow: hidden; box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05); border: 1px solid #E2E8F0;">
            <!-- Header -->
            <tr>
                <td style="background-color: #0F172A; padding: 28px 32px; color: #FFFFFF;">
                    <table width="100%" border="0" cellpadding="0" cellspacing="0">
                        <tr>
                            <td>
                                <div style="font-size: 11px; text-transform: uppercase; letter-spacing: 1px; color: #38BDF8; font-weight: 700; margin-bottom: 4px;">
                                    StayVista Operational Intelligence
                                </div>
                                <h1 style="margin: 0; font-size: 22px; font-weight: 700; color: #FFFFFF; letter-spacing: -0.5px;">
                                    PMS & OTA Automated Reconciliation Report
                                </h1>
                            </td>
                            <td align="right" style="vertical-align: top;">
                                <div style="font-size: 12px; color: #94A3B8; text-align: right;">Date: {today_str}</div>
                                <div style="font-size: 11px; color: #64748B; text-align: right; margin-top: 4px;">Batch: {mis_data.batch_id[:8]}...</div>
                            </td>
                        </tr>
                    </table>
                </td>
            </tr>

            <!-- Notice Banner -->
            <tr>
                <td style="padding: 16px 32px; background-color: #FEF2F2; border-bottom: 1px solid #FEE2E2;">
                    <div style="font-size: 13px; color: #991B1B; font-weight: 500;">
                        <strong>Action Required:</strong> {pending_count} bookings are marked as <strong>Cancellation Pending</strong>. Please review the attached multi-sheet Excel report and execute channel cancellations immediately.
                    </div>
                </td>
            </tr>

            <!-- Content Area -->
            <tr>
                <td style="padding: 28px 32px;">
                    <!-- KPI Metric Cards -->
                    <div style="margin-bottom: 24px;">
                        <h2 style="font-size: 15px; font-weight: 700; color: #0F172A; margin: 0 0 14px 0; text-transform: uppercase; letter-spacing: 0.5px;">
                            Executive Metrics Overview
                        </h2>
                        <table width="100%" border="0" cellpadding="0" cellspacing="0">
                            <tr>
                                <td width="19%" style="background-color: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 6px; padding: 14px; text-align: center;">
                                    <div style="font-size: 11px; color: #64748B; font-weight: 600; text-transform: uppercase;">Total SU</div>
                                    <div style="font-size: 22px; font-weight: 800; color: #0F172A; margin-top: 4px;">{p.total_bookings:,}</div>
                                </td>
                                <td width="2%"></td>
                                <td width="19%" style="background-color: #F0FDF4; border: 1px solid #BBF7D0; border-radius: 6px; padding: 14px; text-align: center;">
                                    <div style="font-size: 11px; color: #166534; font-weight: 600; text-transform: uppercase;">Matched</div>
                                    <div style="font-size: 22px; font-weight: 800; color: #15803D; margin-top: 4px;">{p.matched:,}</div>
                                </td>
                                <td width="2%"></td>
                                <td width="19%" style="background-color: #FEF3C7; border: 1px solid #FDE68A; border-radius: 6px; padding: 14px; text-align: center;">
                                    <div style="font-size: 11px; color: #92400E; font-weight: 600; text-transform: uppercase;">Mismatched</div>
                                    <div style="font-size: 22px; font-weight: 800; color: #B45309; margin-top: 4px;">{p.mismatched:,}</div>
                                </td>
                                <td width="2%"></td>
                                <td width="19%" style="background-color: #F8FAFC; border: 1px solid #E2E8F0; border-radius: 6px; padding: 14px; text-align: center;">
                                    <div style="font-size: 11px; color: #64748B; font-weight: 600; text-transform: uppercase;">Missing in PMS</div>
                                    <div style="font-size: 22px; font-weight: 800; color: #475569; margin-top: 4px;">{p.missing:,}</div>
                                </td>
                                <td width="2%"></td>
                                <td width="19%" style="background-color: #FEF2F2; border: 1px solid #FECACA; border-radius: 6px; padding: 14px; text-align: center;">
                                    <div style="font-size: 11px; color: #991B1B; font-weight: 600; text-transform: uppercase;">Cancel Pending</div>
                                    <div style="font-size: 22px; font-weight: 800; color: #DC2626; margin-top: 4px;">{p.cancellation_pending:,}</div>
                                </td>
                            </tr>
                        </table>
                    </div>

                    <!-- Discrepancy Breakdown Table -->
                    <div style="margin-bottom: 28px;">
                        <h2 style="font-size: 15px; font-weight: 700; color: #0F172A; margin: 0 0 12px 0; text-transform: uppercase; letter-spacing: 0.5px;">
                            Discrepancy Breakdown by Tab
                        </h2>
                        <table width="100%" border="0" cellpadding="0" cellspacing="0" style="border: 1px solid #E2E8F0; border-radius: 6px; border-collapse: separate; overflow: hidden;">
                            <tr style="background-color: #F8FAFC; font-size: 12px; color: #475569; font-weight: 600; text-transform: uppercase; border-bottom: 1px solid #E2E8F0;">
                                <th align="left" style="padding: 10px 14px;">Operational Category</th>
                                <th align="center" style="padding: 10px 14px;">Record Count</th>
                                <th align="left" style="padding: 10px 14px;">Recommended Next Step</th>
                            </tr>
                            <tr style="border-bottom: 1px solid #E2E8F0; font-size: 13px;">
                                <td style="padding: 10px 14px; font-weight: 600; color: #DC2626;">Cancellation Pending</td>
                                <td align="center" style="padding: 10px 14px; font-weight: 700; color: #DC2626;">{mis_data.tab_counts.get('cancellation_pending', 0)}</td>
                                <td style="padding: 10px 14px; color: #475569;">Mark cancelled in PMS & Extranet</td>
                            </tr>
                            <tr style="border-bottom: 1px solid #E2E8F0; font-size: 13px;">
                                <td style="padding: 10px 14px; font-weight: 600; color: #D97706;">Confirmed in SU / Cancelled in PMS</td>
                                <td align="center" style="padding: 10px 14px; font-weight: 700; color: #D97706;">{mis_data.tab_counts.get('confirmed_su_cancelled_pms', 0)}</td>
                                <td style="padding: 10px 14px; color: #475569;">Verify guest check-in & channel commission</td>
                            </tr>
                            <tr style="border-bottom: 1px solid #E2E8F0; font-size: 13px;">
                                <td style="padding: 10px 14px; font-weight: 600; color: #2563EB;">Missing in PMS</td>
                                <td align="center" style="padding: 10px 14px; font-weight: 700; color: #2563EB;">{mis_data.tab_counts.get('missing_in_pms', 0)}</td>
                                <td style="padding: 10px 14px; color: #475569;">Inject valid reservation into PMS</td>
                            </tr>
                            <tr style="border-bottom: 1px solid #E2E8F0; font-size: 13px;">
                                <td style="padding: 10px 14px; font-weight: 600; color: #475569;">Missing in SU (PMS-only)</td>
                                <td align="center" style="padding: 10px 14px; font-weight: 700; color: #475569;">{mis_data.tab_counts.get('missing_in_su', 0)}</td>
                                <td style="padding: 10px 14px; color: #475569;">Audit manual PMS creation entries</td>
                            </tr>
                            <tr style="font-size: 13px;">
                                <td style="padding: 10px 14px; font-weight: 600; color: #059669;">Base vs Query Mismatch</td>
                                <td align="center" style="padding: 10px 14px; font-weight: 700; color: #059669;">{mis_data.tab_counts.get('base_vs_query_mismatch', 0)}</td>
                                <td style="padding: 10px 14px; color: #475569;">Re-sync PMS query report replica</td>
                            </tr>
                        </table>
                    </div>

                    <!-- Priority Action Items (Preview) -->
                    <div style="margin-bottom: 24px;">
                        <div style="display: flex; justify-content: space-between; align-items: baseline; margin-bottom: 12px;">
                            <h2 style="font-size: 15px; font-weight: 700; color: #0F172A; margin: 0; text-transform: uppercase; letter-spacing: 0.5px;">
                                Priority Action Items (Top Cancellation Pending)
                            </h2>
                        </div>
                        <div style="overflow-x: auto; border: 1px solid #E2E8F0; border-radius: 6px;">
                            <table width="100%" border="0" cellpadding="0" cellspacing="0" style="border-collapse: collapse;">
                                <tr style="background-color: #F8FAFC; font-size: 11px; color: #475569; font-weight: 700; text-transform: uppercase; border-bottom: 1px solid #E2E8F0;">
                                    <th align="left" style="padding: 10px;">Res ID</th>
                                    <th align="left" style="padding: 10px;">Vendor ID</th>
                                    <th align="left" style="padding: 10px;">Guest</th>
                                    <th align="left" style="padding: 10px;">Channel</th>
                                    <th align="left" style="padding: 10px;">Property</th>
                                    <th align="left" style="padding: 10px;">Assigned RM</th>
                                    <th align="left" style="padding: 10px;">Check-In</th>
                                    <th align="left" style="padding: 10px;">OTA Extranet</th>
                                </tr>
                                {action_rows_html}
                            </table>
                        </div>
                    </div>

                    <!-- Attachment Callout -->
                    <div style="background-color: #F8FAFC; border: 1px dashed #CBD5E1; border-radius: 6px; padding: 16px; margin-top: 24px; text-align: center;">
                        <div style="font-size: 13px; font-weight: 600; color: #1E293B;">
                            📎 Complete Multi-Sheet Excel Workbook Attached
                        </div>
                        <div style="font-size: 12px; color: #64748B; margin-top: 4px;">
                            Contains complete datasets for all 8 categories: Cancellation Pending, Status Mismatches, Missing Records, and Representative Assignments.
                        </div>
                    </div>
                </td>
            </tr>

            <!-- Footer -->
            <tr>
                <td style="background-color: #F1F5F9; padding: 18px 32px; border-top: 1px solid #E2E8F0; font-size: 11px; color: #64748B; text-align: center;">
                    This automated reconciliation alert was generated by the <strong>StayVista PMS & OTA Reconciliation Engine</strong>.<br>
                    Run timestamp: {now_str} | Batch ID: {mis_data.batch_id}
                </td>
            </tr>
        </table>
    </body>
    </html>
    """
    return html


def parse_recipient_list(raw_recipients: Optional[Any]) -> List[str]:
    """Parse comma/semicolon-separated string or list of recipient email addresses."""
    if not raw_recipients:
        return []
    
    if isinstance(raw_recipients, (list, tuple, set)):
        items = raw_recipients
    else:
        # String separated by comma or semicolon
        items = str(raw_recipients).replace(";", ",").split(",")

    valid = []
    for it in items:
        clean = str(it).strip()
        if clean and "@" in clean and "." in clean.split("@")[-1]:
            valid.append(clean)
    return valid


def save_dry_run_email(
    subject: str,
    recipients: List[str],
    html_body: str,
    excel_bytes: bytes,
    batch_id: str,
    reason: str = "Dry run mode / SMTP credentials not configured"
) -> Dict[str, Any]:
    """
    When SMTP is not configured or in office dry-run mode, safely archive
    the email and report without failing and without using any personal email.
    """
    archive_dir = settings.BASE_DIR / "archives" / "emails"
    os.makedirs(archive_dir, exist_ok=True)

    today_str = date.today().strftime("%Y-%m-%d")
    html_path = archive_dir / f"email_{today_str}_{batch_id[:8]}.html"
    excel_path = archive_dir / f"StayVista_Reconciliation_Report_{today_str}_{batch_id[:8]}.xlsx"

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html_body)

    with open(excel_path, "wb") as f:
        f.write(excel_bytes)

    return {
        "status": "dry_run_saved",
        "message": f"Office email report generated and archived successfully ({reason}).",
        "recipients": recipients,
        "subject": subject,
        "html_preview_path": str(html_path),
        "excel_attachment_path": str(excel_path),
        "attachment_size_kb": round(len(excel_bytes) / 1024, 1),
    }


def send_reconciliation_email(
    mis_data: MISDashboardData,
    recipient_emails: Optional[List[str]] = None,
    custom_subject: Optional[str] = None,
    dry_run_if_no_smtp: bool = True
) -> Dict[str, Any]:
    """
    Send the comprehensive StayVista MIS Reconciliation Report and Excel attachment
    to the configured team/office recipient list.
    
    - Uses free standard SMTP (Office 365, Google Workspace, or custom corporate SMTP).
    - Never uses user's personal email.
    - If SMTP is not yet configured, cleanly archives the report in dry_run mode so no errors occur.
    """
    # 1. Resolve recipients
    recipients = parse_recipient_list(recipient_emails)
    if not recipients:
        recipients = parse_recipient_list(settings.RECIPIENT_EMAILS)

    # 2. Build Excel Attachment & HTML Body
    excel_bytes = generate_styled_excel_report(mis_data)
    html_body = build_email_html_body(mis_data)

    today_str = date.today().strftime("%Y-%m-%d")
    p = mis_data.primary_metrics
    subject = custom_subject or (
        f"[StayVista MIS Alert] Reconciliation Report ({today_str}) - "
        f"{p.cancellation_pending} Pending Cancellations | {p.matched} Matched"
    )

    # 3. Check SMTP configuration
    smtp_host = (settings.SMTP_HOST or "").strip()
    smtp_port = int(settings.SMTP_PORT or 587)
    smtp_user = (settings.SMTP_USER or "").strip()
    smtp_pass = (settings.SMTP_PASSWORD or "").strip()
    from_name = (settings.SMTP_FROM_NAME or "StayVista Reconciliation Engine").strip()

    if not smtp_user or not smtp_pass:
        if dry_run_if_no_smtp:
            return save_dry_run_email(
                subject=subject,
                recipients=recipients,
                html_body=html_body,
                excel_bytes=excel_bytes,
                batch_id=mis_data.batch_id,
                reason="Office SMTP credentials (SMTP_USER/SMTP_PASSWORD) are not set in .env"
            )
        else:
            raise ValueError(
                "Office email credentials not configured! Please configure SMTP_USER and SMTP_PASSWORD in .env"
            )

    if not recipients:
        return {
            "status": "warning",
            "message": "No recipient emails provided. Please specify RECIPIENT_EMAILS in .env or via API.",
            "recipients": [],
            "subject": subject
        }

    # 4. Construct Multipart Email Message
    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = f"{from_name} <{smtp_user}>"
    msg["To"] = ", ".join(recipients)

    # HTML Part
    msg_html = MIMEText(html_body, "html", "utf-8")
    msg.attach(msg_html)

    # Excel Attachment Part
    excel_filename = f"StayVista_Reconciliation_Report_{today_str}_{mis_data.batch_id[:8]}.xlsx"
    part_excel = MIMEApplication(excel_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part_excel.add_header("Content-Disposition", "attachment", filename=excel_filename)
    msg.attach(part_excel)

    # 5. Connect to SMTP server and transmit
    try:
        # Support TLS (port 587) or SSL (port 465)
        if smtp_port == 465:
            server = smtplib.SMTP_SSL(smtp_host, smtp_port, timeout=30)
        else:
            server = smtplib.SMTP(smtp_host, smtp_port, timeout=30)
            server.ehlo()
            if settings.SMTP_USE_TLS:
                server.starttls()
                server.ehlo()

        server.login(smtp_user, smtp_pass)
        server.sendmail(smtp_user, recipients, msg.as_string())
        server.quit()

        # Archive sent copy for audit
        save_dry_run_email(
            subject=subject,
            recipients=recipients,
            html_body=html_body,
            excel_bytes=excel_bytes,
            batch_id=mis_data.batch_id,
            reason="Sent successfully via SMTP"
        )

        return {
            "status": "success",
            "message": f"Successfully sent reconciliation report to {len(recipients)} recipient(s).",
            "recipients": recipients,
            "subject": subject,
            "batch_id": mis_data.batch_id,
            "attachment_name": excel_filename,
            "attachment_size_kb": round(len(excel_bytes) / 1024, 1),
            "timestamp": datetime.now().isoformat()
        }
    except Exception as e:
        # Also archive local backup so work is never lost
        save_dry_run_email(
            subject=subject,
            recipients=recipients,
            html_body=html_body,
            excel_bytes=excel_bytes,
            batch_id=mis_data.batch_id,
            reason=f"SMTP transmission error: {str(e)}"
        )
        return {
            "status": "error",
            "message": f"Failed to send email via SMTP ({smtp_host}:{smtp_port}): {str(e)}",
            "recipients": recipients,
            "subject": subject
        }


def send_test_email(recipient_email: str) -> Dict[str, Any]:
    """Send a quick test email to verify office SMTP credentials without running full reconciliation."""
    smtp_host = (settings.SMTP_HOST or "").strip()
    smtp_port = int(settings.SMTP_PORT or 587)
    smtp_user = (settings.SMTP_USER or "").strip()
    smtp_pass = (settings.SMTP_PASSWORD or "").strip()
    from_name = (settings.SMTP_FROM_NAME or "StayVista Reconciliation Engine").strip()

    if not smtp_user or not smtp_pass:
        return {
            "status": "error",
            "message": "Office SMTP credentials are empty. Please set SMTP_USER and SMTP_PASSWORD in .env."
        }

    msg = MIMEMultipart("alternative")
    msg["Subject"] = "[StayVista] Office SMTP Test Connection Successful"
    msg["From"] = f"{from_name} <{smtp_user}>"
    msg["To"] = recipient_email

    html = f"""
    <div style="font-family: sans-serif; padding: 20px; color: #1E293B;">
        <h2 style="color: #0F172A;">StayVista Automation Test</h2>
        <p>This is a verification email from your office reconciliation engine.</p>
        <p style="color: #10B981; font-weight: bold;">✔ Office SMTP connection was established and verified successfully!</p>
        <p style="font-size: 12px; color: #64748B;">Server: {smtp_host}:{smtp_port} | Sender: {smtp_user}</p>
    </div>
    """
    msg.attach(MIMEText(html, "html", "utf-8"))

    try:
        if smtp_port == 465:
            server = smtplib.SMTP_SSL(smtp_host, smtp_port, timeout=20)
        else:
            server = smtplib.SMTP(smtp_host, smtp_port, timeout=20)
            server.ehlo()
            if settings.SMTP_USE_TLS:
                server.starttls()
                server.ehlo()
        server.login(smtp_user, smtp_pass)
        server.sendmail(smtp_user, [recipient_email], msg.as_string())
        server.quit()
        return {
            "status": "success",
            "message": f"Test email sent successfully to {recipient_email}."
        }
    except Exception as e:
        return {
            "status": "error",
            "message": f"SMTP test failed: {str(e)}"
        }
